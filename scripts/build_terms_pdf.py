"""Export the four published policy pages as PDF, DOCX and Markdown.

Run export-policy-source.cjs first. This builder archives the preceding bot
agreement and updates its manifest plus the Go consent constants together.
Use a new VERSION for any change to an agreement already committed at HEAD.
"""
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from xml.sax.saxutils import escape
import hashlib
import json
import re
import shutil
import subprocess

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
VERSION = "2026-10-16"
BASE = "https://pairtalk.online"
source_file = ROOT / "output/policies/source.json"
source = json.loads(source_file.read_text(encoding="utf-8"))
assert [page["route"] for page in source] == ["/terms", "/privacy", "/community-guidelines", "/refund-policy"]
manifest_path = ROOT / "server/assets/terms-manifest.json"
committed = subprocess.run(["git", "show", "HEAD:server/assets/terms-manifest.json"], cwd=ROOT, capture_output=True, text=True, check=True)
assert json.loads(committed.stdout)["version"] != VERSION, "Choose a new version for an already committed agreement."

def plain(text):
    return " ".join(text.split()).translate(str.maketrans({"\u2013": "-", "\u2014": "-", "\u2011": "-", "\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u00a0": " "}))

class Paragraphs(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.items, self.current, self.link = [], None, None
    def handle_starttag(self, tag, attrs):
        if tag in ("p", "li"):
            self.current = []
        if tag == "a":
            self.link = dict(attrs).get("href")
        if tag == "br" and self.current is not None:
            self.current.append(" ")
    def handle_data(self, text):
        if self.current is not None:
            self.current.append(text)
    def handle_endtag(self, tag):
        if tag == "a" and self.link and self.current is not None:
            self.current.append(" (" + (BASE + self.link if self.link.startswith("/") else self.link) + ")")
            self.link = None
        if tag in ("p", "li") and self.current is not None:
            self.items.append(plain("".join(self.current)))
            self.current = None

titles = {"/terms": "Terms of use", "/privacy": "Privacy policy", "/community-guidelines": "Community guidelines", "/refund-policy": "Refund policy"}
for page in source:
    page["title"] = titles[page["route"]]
    for section in page["sections"]:
        parser = Paragraphs(); parser.feed(section.pop("html")); section["paragraphs"] = parser.items
        assert parser.items, "Every policy section must contain readable paragraphs."

out_pdf = ROOT / "output/pdf/pairtalk-policies-2026-10-16.pdf"
out_docx = ROOT / "output/doc/pairtalk-policies-2026-10-16.docx"
out_pdf.parent.mkdir(parents=True, exist_ok=True); out_docx.parent.mkdir(parents=True, exist_ok=True)
md_dir = ROOT / "docs/policies" / VERSION
md_dir.mkdir(parents=True, exist_ok=True)
ink = colors.HexColor("#172923")
body = ParagraphStyle("Body", fontName="Helvetica", fontSize=11, leading=16, textColor=ink, spaceAfter=10)
heading = ParagraphStyle("Heading", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=18, spaceBefore=12, spaceAfter=8, keepWithNext=True)
title = ParagraphStyle("Title", parent=body, fontName="Helvetica-Bold", fontSize=25, leading=30, spaceAfter=15)
small = ParagraphStyle("Small", parent=body, fontSize=9, leading=13, textColor=colors.HexColor("#52635c"))

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5); section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = document.styles[name]
    style.font.name = "Arial"; style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    for border in style.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
normal = document.styles["Normal"]
normal.font.size = Pt(11); normal.paragraph_format.line_spacing = 1.15; normal.paragraph_format.space_after = Pt(8)
document.styles["Title"].font.size = Pt(26)
document.styles["Heading 1"].font.size = Pt(21)
document.styles["Heading 2"].font.size = Pt(13)
for name in ("Heading 1", "Heading 2"):
    document.styles[name].paragraph_format.keep_with_next = True
document.core_properties.title = "PairTalk service policies"
document.core_properties.author = "PairTalk"
document.core_properties.subject = "Terms privacy community guidelines and refunds"
document.core_properties.created = document.core_properties.modified = datetime(2026, 10, 16, tzinfo=timezone.utc)

intro = "Read these policies before using PairTalk. They explain purchases and public offers, your privacy choices, respectful speaking practice and refunds. Agree to this version in the bot before registered practice; recording always needs separate agreement with your partner."
document.add_paragraph("PairTalk service policies", "Title")
document.add_paragraph(f"Version {VERSION} | Effective 16 October 2026")
document.add_paragraph(intro)
story = [Paragraph("PairTalk service policies", title), Paragraph(f"Version {VERSION} | Effective 16 October 2026", small), Paragraph(escape(intro), body)]
for index, page in enumerate(source):
    if index:
        story.append(PageBreak()); document.add_page_break()
    story.extend([Spacer(1, 12), Paragraph(escape(page["title"]), title), Paragraph(escape(plain(page["intro"])), body), Paragraph(f'Published policy: {BASE}{page["route"]}', small)])
    document.add_heading(page["title"], 1); document.add_paragraph(plain(page["intro"]))
    document.add_paragraph(f'Published policy: {BASE}{page["route"]}')
    markdown = [f'# {page["title"]}', "", f'Version {VERSION}. Effective 16 October 2026.', "", plain(page["intro"]), "", f'Published page: {BASE}{page["route"]}', ""]
    for item in page["sections"]:
        story.append(Paragraph(escape(item["title"]), heading))
        # Word heading titles contain words and spaces only.
        document.add_heading(re.sub(r"[^\w\s]", "", item["title"]), 2)
        markdown.extend([f'## {item["title"]}', ""])
        for paragraph in item["paragraphs"]:
            story.append(Paragraph(escape(paragraph), body)); document.add_paragraph(paragraph)
            markdown.extend([paragraph, ""])
    (md_dir / (page["route"].strip("/") + ".md")).write_text("\n".join(markdown), encoding="utf-8")

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont("Helvetica", 8); canvas.setFillColor(colors.HexColor("#52635c"))
    canvas.drawString(58, 26, f"PairTalk | Policy version {VERSION}")
    canvas.drawRightString(554, 26, str(doc.page)); canvas.restoreState()

SimpleDocTemplate(str(out_pdf), pagesize=(612, 792), leftMargin=58, rightMargin=58, topMargin=48, bottomMargin=44, title="PairTalk service policies", author="PairTalk", pageCompression=1, invariant=1).build(story, onFirstPage=footer, onLaterPages=footer)
footer_p = document.sections[0].footer.paragraphs[0]
footer_p.add_run(f"PairTalk | Policy version {VERSION} | ")
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer_p._p.append(field)
document.save(out_docx)

text = plain(" ".join(re.sub(r"PairTalk\s*\|\s*Policy version " + VERSION + r"\s*\d+", "", page.extract_text()) for page in PdfReader(out_pdf).pages))
for page in source:
    for item in page["sections"]:
        for paragraph in item["paragraphs"]:
            assert plain(paragraph) in text, f'Missing PDF text: {item["title"]}'
doc_text = " ".join(p.text for p in Document(out_docx).paragraphs)
for page in source:
    for item in page["sections"]:
        for paragraph in item["paragraphs"]:
            assert paragraph in doc_text, f'Missing DOCX text: {item["title"]}'

asset = ROOT / "server/assets/pairtalk-terms-of-use.pdf"
previous = json.loads(manifest_path.read_text(encoding="utf-8"))
if previous["version"] != VERSION:
    archive = asset.parent / "terms" / previous["version"]
    archive.mkdir(parents=True, exist_ok=True)
    for old, name in ((asset, "pairtalk-terms-of-use.pdf"), (manifest_path, "terms-manifest.json")):
        existing = archive / name
        if existing.exists():
            assert existing.read_bytes() == old.read_bytes(), "An accepted archived policy cannot be overwritten."
        else:
            shutil.copy2(old, existing)
shutil.copy2(out_pdf, asset)
published = ROOT / "landing/public/policies"
published.mkdir(parents=True, exist_ok=True)
shutil.copy2(out_pdf, published / out_pdf.name)
# The owner deferred public DOCX delivery until native page rendering is available.
# Keep the editable artifact local; do not publish an unrendered document.
manifest = {"version": VERSION, "filename": asset.name, "sha256": hashlib.sha256(asset.read_bytes()).hexdigest(), "source": "landing/src/content.tsx", "termsUrl": BASE + "/terms", "privacyUrl": BASE + "/privacy"}
manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
(ROOT / "gateway/internal/auth/terms.go").write_text(f'package auth\n\n// Matches the deployed PDF manifest. Update using scripts/build_terms_pdf.py.\nconst TermsVersion = "{VERSION}"\nconst TermsDocumentSHA256 = "{manifest["sha256"]}"\n', encoding="utf-8")
print(f'Created PDF ({len(PdfReader(out_pdf).pages)} pages), editable DOCX and four Markdown policies; version {VERSION}.')
